import os
import openpyxl
from decimal import Decimal

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.db.models import Sum
from django.utils import timezone

from apps.accounting.models import JournalEntry, JournalItem, Account, AccountingFiscalYear
from apps.purchases.models import GoodsReceivedNote, PurchaseReturn, PurchaseReturnItem, Supplier
from apps.sales.models import SalesEstimate, SalesReturn, SalesReturnItem, Customer
from apps.inventory.models import Product, BranchStock, StockMovementLog
from apps.branches.models import Branch


class Command(BaseCommand):
    help = "Production command executing Steps 1 to 4 of the Hisaav Legacy Migration."

    def add_arguments(self, parser):
        parser.add_argument(
            '--data-dir',
            type=str,
            default=r"C:\Users\Vivek Mani Upadhyaya\Downloads\Mobile soft Data",
            help="Path to folder containing legacy Hisaav Excel files."
        )
        parser.add_argument(
            '--step',
            type=int,
            choices=[1, 2, 3, 4],
            default=None,
            help="Execute a specific step (1=Cancellations, 2=Purchase Returns, 3=Sales Returns, 4=Reconcile & Lock). Default: all."
        )
        parser.add_argument(
            '--dry-run',
            action='store_true',
            help="Simulate execution without committing changes."
        )

    def handle(self, *args, **options):
        data_dir = options['data_dir']
        step_filter = options['step']
        dry_run = options['dry_run']

        self.stdout.write(self.style.SUCCESS("=" * 85))
        self.stdout.write(self.style.SUCCESS("HISAAV LEGACY MIGRATION PIPELINE"))
        self.stdout.write(self.style.SUCCESS(f"Data Directory : {data_dir}"))
        self.stdout.write(self.style.SUCCESS(f"Mode           : {'DRY RUN (NO COMMIT)' if dry_run else 'LIVE COMMIT'}"))
        self.stdout.write(self.style.SUCCESS("=" * 85))

        if not os.path.exists(data_dir):
            raise CommandError(f"Data directory does not exist: {data_dir}")

        branch = Branch.objects.filter(code='BR-MAIN-01').first() or Branch.objects.first()
        if not branch:
            raise CommandError("No Branch found in the database. Ensure basic branch setup exists.")

        try:
            with transaction.atomic():
                # Step 1: Purchase Cancellations
                if step_filter is None or step_filter == 1:
                    self.run_step_1_cancellations(data_dir)

                # Step 2: Purchase Returns (Debit Notes)
                if step_filter is None or step_filter == 2:
                    self.run_step_2_purchase_returns(data_dir, branch)

                # Step 3: Sales Returns (Credit Notes)
                if step_filter is None or step_filter == 3:
                    self.run_step_3_sales_returns(data_dir, branch)

                # Step 4: Parity Audit & Fiscal Year Lock
                if step_filter is None or step_filter == 4:
                    self.run_step_4_reconciliation_and_locks()

                if dry_run:
                    self.stdout.write(self.style.WARNING("\n[DRY-RUN] Simulation successful. Rolling back all database changes."))
                    transaction.set_rollback(True)
                else:
                    self.stdout.write(self.style.SUCCESS("\n[SUCCESS] Pipeline executed and committed successfully."))

        except Exception as e:
            self.stderr.write(self.style.ERROR(f"\n[FATAL ERROR] Pipeline aborted: {str(e)}"))
            raise

    # -------------------------------------------------------------------------
    # STEP 1: PURCHASE CANCELLATIONS
    # -------------------------------------------------------------------------
    def run_step_1_cancellations(self, data_dir):
        self.stdout.write("\n" + "-" * 75)
        self.stdout.write(self.style.MIGRATE_HEADING("STEP 1: Processing Purchase Cancellations (Bill Cancel Report.xlsx)"))
        self.stdout.write("-" * 75)

        file_path = os.path.join(data_dir, "Bill Cancel Report.xlsx")
        if not os.path.exists(file_path):
            self.stdout.write(self.style.WARNING(f"File not found: {file_path}. Skipping Step 1."))
            return

        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb.active

        cancelled_bills = set()
        for r in range(8, ws.max_row + 1):
            bill_no = str(ws.cell(r, 2).value or '').strip()
            if bill_no and bill_no not in ['None', '', '-']:
                cancelled_bills.add(bill_no)

        self.stdout.write(f"Found {len(cancelled_bills)} cancellation entries in Excel.")

        cancelled_count = 0
        reversed_amount = Decimal('0.00')

        for bill_no in cancelled_bills:
            grns = GoodsReceivedNote.objects.filter(bill_number=bill_no)
            for grn in grns:
                if grn.status != 'CANCELLED':
                    grn.status = 'CANCELLED'
                    grn.due_amount = Decimal('0.00')
                    grn.save(update_fields=['status', 'due_amount', 'updated_at'])

                    # Cancel associated journal entry
                    jes = JournalEntry.objects.filter(reference_document=grn.grn_number, status='POSTED')
                    for je in jes:
                        je.status = 'CANCELLED'
                        je.narration = f"[MIGRATION VOID] {je.narration}"
                        je.save(update_fields=['status', 'narration', 'updated_at'])

                    # Stock reversal log
                    StockMovementLog.objects.get_or_create(
                        reference_document=f"VOID-{grn.grn_number}",
                        defaults={
                            'movement_type': 'ADJUSTMENT_SUB',
                            'branch': grn.branch,
                            'notes': f"Migration stock reversal for cancelled purchase bill #{bill_no}"
                        }
                    )
                    cancelled_count += 1
                    reversed_amount += grn.gross_amount

        self.stdout.write(self.style.SUCCESS(
            f"Step 1 Complete: {cancelled_count} GRNs cancelled. Reversed: Rs. {reversed_amount:,.2f}"
        ))

    # -------------------------------------------------------------------------
    # STEP 2: PURCHASE RETURNS / DEBIT NOTES
    # -------------------------------------------------------------------------
    def run_step_2_purchase_returns(self, data_dir, branch):
        self.stdout.write("\n" + "-" * 75)
        self.stdout.write(self.style.MIGRATE_HEADING("STEP 2: Processing Purchase Returns / Debit Notes"))
        self.stdout.write("-" * 75)

        file_path = os.path.join(data_dir, "Purchase Return Report9_18_2026.xlsx")
        if not os.path.exists(file_path):
            self.stdout.write(self.style.WARNING(f"File not found: {file_path}. Skipping Step 2."))
            return

        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb.active

        cash_account = Account.objects.filter(code='1110-BR-MAIN-01').first() or Account.objects.filter(account_type='ASSET', name__icontains='Cash').first()
        inventory_account = Account.objects.filter(code='1310-BR-MAIN-01').first() or Account.objects.filter(account_type='ASSET', name__icontains='Inventory').first()
        historical_product, _ = Product.objects.get_or_create(
            name="Mobile (Historical)",
            defaults={'sku': 'HIST-MOB-001', 'is_active': True}
        )

        return_rows = []
        for r in range(11, ws.max_row + 1):
            date_bs = str(ws.cell(r, 1).value or '').strip()
            sup_name = str(ws.cell(r, 4).value or '').strip()
            pan = str(ws.cell(r, 5).value or '').strip()
            gross_val = ws.cell(r, 9).value
            taxable_val = ws.cell(r, 11).value
            vat_val = ws.cell(r, 12).value

            if gross_val and sup_name and date_bs not in ['None', '', '-']:
                return_rows.append({
                    'date_bs': date_bs,
                    'supplier_name': sup_name,
                    'pan': pan,
                    'gross': Decimal(str(gross_val)),
                    'taxable': Decimal(str(taxable_val or gross_val)),
                    'vat': Decimal(str(vat_val or '0.00'))
                })

        created_dn_count = 0
        total_dn_amount = Decimal('0.00')

        for idx, item in enumerate(return_rows, start=1):
            dn_number = f"DN-BR-MAIN-01-{idx:06d}"
            pr_number = f"PR-BR-MAIN-01-{idx:06d}"

            supplier = Supplier.objects.filter(pan_number=item['pan']).first() or Supplier.objects.filter(company_name__icontains=item['supplier_name']).first()
            if not supplier:
                supplier = Supplier.objects.create(
                    company_name=item['supplier_name'],
                    pan_number=item['pan'],
                    is_active=True
                )
            elif item['pan'] and not supplier.pan_number:
                supplier.pan_number = item['pan']
                supplier.save(update_fields=['pan_number'])

            pr, created_pr = PurchaseReturn.objects.get_or_create(
                return_number=pr_number,
                defaults={
                    'branch': branch,
                    'supplier': supplier,
                    'status': 'CONFIRMED',
                    'refund_mode': 'CASH_REFUND',
                    'taxable_amount': item['taxable'],
                    'vat_amount': item['vat'],
                    'total_amount': item['gross'],
                    'reason': f"Hisaav Purchase Return - {item['date_bs']}"
                }
            )

            if created_pr:
                PurchaseReturnItem.objects.get_or_create(
                    purchase_return=pr,
                    product=historical_product,
                    defaults={
                        'quantity': Decimal('1.00'),
                        'unit_price': item['taxable'],
                        'taxable_amount': item['taxable'],
                        'vat_amount': item['vat'],
                        'total_amount': item['gross']
                    }
                )

            # Debit Note Voucher
            je, created_je = JournalEntry.objects.get_or_create(
                voucher_number=dn_number,
                defaults={
                    'voucher_type': 'DEBIT_NOTE',
                    'branch': branch,
                    'reference_document': pr_number,
                    'narration': f"Debit Note for Purchase Return {pr_number} - {item['supplier_name']}",
                    'entry_date': timezone.now().date(),
                    'entry_date_bs': item['date_bs'],
                    'fiscal_year': '2083/84',
                    'total_debit': item['gross'],
                    'total_credit': item['gross'],
                    'status': 'POSTED'
                }
            )

            if created_je:
                JournalItem.objects.create(
                    journal_entry=je,
                    account=cash_account,
                    debit_amount=item['gross'],
                    credit_amount=Decimal('0.00'),
                    line_narration=f"Cash refund from supplier {item['supplier_name']}"
                )
                JournalItem.objects.create(
                    journal_entry=je,
                    account=inventory_account,
                    debit_amount=Decimal('0.00'),
                    credit_amount=item['gross'],
                    line_narration="Inventory reversal for purchase return"
                )
                created_dn_count += 1
                total_dn_amount += item['gross']

        self.stdout.write(self.style.SUCCESS(
            f"Step 2 Complete: {created_dn_count} Debit Notes created. Gross Reversal: Rs. {total_dn_amount:,.2f}"
        ))

    # -------------------------------------------------------------------------
    # STEP 3: SALES RETURNS / CREDIT NOTES
    # -------------------------------------------------------------------------
    def run_step_3_sales_returns(self, data_dir, branch):
        self.stdout.write("\n" + "-" * 75)
        self.stdout.write(self.style.MIGRATE_HEADING("STEP 3: Processing Sales Returns / Credit Notes"))
        self.stdout.write("-" * 75)

        sr_number = "SR-BR-MAIN-01-000001"
        cn_number = "CN-BR-MAIN-01-000001"

        if SalesReturn.objects.filter(return_number=sr_number).exists():
            self.stdout.write(self.style.WARNING("Step 3 Sales Return already exists. Skipping duplicate."))
            return

        taxable = Decimal('156637.17')
        vat = Decimal('20362.83')
        gross = Decimal('177000.00')

        customer = Customer.objects.filter(pan_number='304427312').first() or Customer.objects.first()
        product = Product.objects.filter(name__icontains='I Phone 16 Pro 256').first() or Product.objects.first()
        estimate = SalesEstimate.objects.filter(grand_total=gross).first()

        sr = SalesReturn.objects.create(
            branch=branch,
            customer=customer,
            return_number=sr_number,
            sales_estimate=estimate,
            refund_mode='CASH',
            taxable_amount=taxable,
            vat_amount=vat,
            total_amount=gross,
            reason="Customer sales return: I Phone 16 Pro 256"
        )

        SalesReturnItem.objects.create(
            sales_return=sr,
            product=product,
            quantity=Decimal('1.000'),
            unit_price=taxable,
            taxable_amount=taxable,
            vat_amount=vat,
            total_amount=gross
        )

        bs, _ = BranchStock.objects.get_or_create(branch=branch, product=product)
        bs.quantity += Decimal('1.000')
        bs.save(update_fields=['quantity', 'updated_at'])

        StockMovementLog.objects.create(
            movement_type='SALE_RETURN',
            branch=branch,
            product=product,
            quantity=Decimal('1.000'),
            reference_document=sr_number,
            notes="Restocked 1 unit via Sales Return SR-BR-MAIN-01-000001"
        )

        sales_acc = Account.objects.filter(code='4110-BR-MAIN-01').first() or Account.objects.filter(account_type='REVENUE').first()
        vat_acc = Account.objects.filter(code='2210-BR-MAIN-01').first() or Account.objects.filter(name__icontains='Output VAT').first()
        cash_acc = Account.objects.filter(code='1110-BR-MAIN-01').first() or Account.objects.filter(account_type='ASSET', name__icontains='Cash').first()

        je = JournalEntry.objects.create(
            voucher_number=cn_number,
            voucher_type='CREDIT_NOTE',
            branch=branch,
            reference_document=sr_number,
            narration=f"Credit Note for Sales Return {sr_number}",
            entry_date=timezone.now().date(),
            entry_date_bs="2083-05-18",
            fiscal_year="2083/84",
            total_debit=gross,
            total_credit=gross,
            status='POSTED'
        )

        JournalItem.objects.create(
            journal_entry=je, account=sales_acc, debit_amount=taxable, credit_amount=Decimal('0.00'),
            line_narration="Sales revenue reversal on customer return"
        )
        JournalItem.objects.create(
            journal_entry=je, account=vat_acc, debit_amount=vat, credit_amount=Decimal('0.00'),
            line_narration="Output VAT 13% reversal on customer return"
        )
        JournalItem.objects.create(
            journal_entry=je, account=cash_acc, debit_amount=Decimal('0.00'), credit_amount=gross,
            line_narration="Cash refund paid to customer"
        )

        self.stdout.write(self.style.SUCCESS(
            f"Step 3 Complete: Credit Note {cn_number} posted. Gross Refund: Rs. {gross:,.2f}"
        ))

    # -------------------------------------------------------------------------
    # STEP 4: AUDIT & FISCAL LOCK ENFORCEMENT
    # -------------------------------------------------------------------------
    def run_step_4_reconciliation_and_locks(self):
        self.stdout.write("\n" + "-" * 75)
        self.stdout.write(self.style.MIGRATE_HEADING("STEP 4: Parity Audit & Fiscal Lock Enforcement"))
        self.stdout.write("-" * 75)

        posted_items = JournalItem.objects.filter(journal_entry__status='POSTED')
        totals = posted_items.aggregate(total_dr=Sum('debit_amount'), total_cr=Sum('credit_amount'))
        total_dr = totals['total_dr'] or Decimal('0.00')
        total_cr = totals['total_cr'] or Decimal('0.00')
        diff = total_dr - total_cr

        self.stdout.write(f"GL Total Debits : Rs. {total_dr:,.2f}")
        self.stdout.write(f"GL Total Credits: Rs. {total_cr:,.2f}")
        self.stdout.write(f"Internal GL Diff: Rs. {diff:,.2f}")

        if diff != Decimal('0.00'):
            raise CommandError(f"CRITICAL GL IMBALANCE: Dr - Cr delta is Rs. {diff:,.2f}!")

        # Re-close historical fiscal years
        for fy_name in ['2080/81', '2081/82', '2082/83']:
            AccountingFiscalYear.objects.filter(name=fy_name).update(is_closed=True, is_active=False)
            self.stdout.write(f"Enforced Lock: FY {fy_name} -> is_closed=True, is_active=False")

        # Keep current active year open
        AccountingFiscalYear.objects.filter(name='2083/84').update(is_closed=False, is_active=True)
        self.stdout.write(f"Enforced Active: FY 2083/84 -> is_closed=False, is_active=True")

        self.stdout.write(self.style.SUCCESS("Step 4 Complete: Parity verified (0.00) and historical fiscal years locked."))